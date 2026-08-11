from __future__ import annotations



def parse_config_text(raw_text):
    """Parse FortiGate config text without using Flask route handlers."""
    text = raw_text if isinstance(raw_text, str) else str(raw_text or "")
    parser = FortiGateConfigParser(text)
    return parser.parse()



def _normalize_items(items):
    """Return a safe list of dictionaries for diff comparison."""
    if items is None:
        return []
    if isinstance(items, dict):
        return [items]
    if isinstance(items, list):
        return [item for item in items if isinstance(item, dict)]
    return []

def _build_policy_map(items):
    result = {}
    for item in _normalize_items(items):
        pid = str(item.get("policy_id") or item.get("id") or item.get("_edit") or "").strip()
        if pid:
            result[pid] = item
    return result


def _index_named_items(items, name_key="name"):
    result = {}
    for item in _normalize_items(items):
        name = str(item.get(name_key) or item.get("_edit") or item.get("policy_id") or item.get("id") or "").strip()
        if name:
            result[name] = item
    return result


def _simplify_policy_for_compare(policy):
    if not isinstance(policy, dict):
        return {}
    # hit_count·last_used는 런타임 통계라 diff 노이즈지만, status는 config
    # 속성이다 — 정책 활성↔비활성 전환은 검토자가 반드시 봐야 하는 변경인데
    # 무시되고 있었다(v79 감사 증적 자체 재비판에서 발견).
    ignored_keys = {"hit_count", "last_used"}
    return {k: v for k, v in policy.items() if k not in ignored_keys}




def _section_signature(items):
    result = {}
    for item in _normalize_items(items):
        key = str(item.get("name") or item.get("_edit") or item.get("policy_id") or item.get("id") or "").strip()
        if key:
            cleaned = {k: v for k, v in item.items() if k not in {"uuid"}}
            result[key] = cleaned
    return result

def _compute_config_diff(old_data, new_data):
    old_policies = _build_policy_map(old_data.get("firewall_policy") or old_data.get("policies") or [])
    new_policies = _build_policy_map(new_data.get("firewall_policy") or new_data.get("policies") or [])

    added_policies = []
    removed_policies = []
    changed_policies = []

    for pid, policy in new_policies.items():
        if pid not in old_policies:
            added_policies.append(policy)
        else:
            if _simplify_policy_for_compare(old_policies[pid]) != _simplify_policy_for_compare(policy):
                changed_policies.append({
                    # added/removed는 정책 객체의 int ID를 그대로 내는데 여기만
                    # 맵 키(str)를 내고 있었다(감사 C6). 타입을 맞춘다.
                    "policy_id": int(pid) if str(pid).isdigit() else pid,
                    "name": policy.get("name") or old_policies[pid].get("name") or "",
                    "before": old_policies[pid],
                    "after": policy,
                })

    for pid, policy in old_policies.items():
        if pid not in new_policies:
            removed_policies.append(policy)

    object_sections = [
        "firewall_address",
        "firewall_addrgrp",
        "firewall_proxy_address",
        "firewall_proxy_addrgrp",
        "firewall_service_custom",
        "firewall_service_group",
        "system_interface",
    ]

    added_objects = []
    removed_objects = []

    for section in object_sections:
        old_map = _index_named_items(old_data.get(section) or [])
        new_map = _index_named_items(new_data.get(section) or [])

        for name, item in new_map.items():
            if name not in old_map:
                added_objects.append({
                    "section": section,
                    "name": name,
                    "item": item,
                })

        for name, item in old_map.items():
            if name not in new_map:
                removed_objects.append({
                    "section": section,
                    "name": name,
                    "item": item,
                })


    known_sections = {
        "firewall_policy",
        "policies",
        "firewall_address",
        "firewall_addrgrp",
        "firewall_proxy_address",
        "firewall_proxy_addrgrp",
        "firewall_service_custom",
        "firewall_service_group",
        "system_interface",
    }
    other_changes = []
    all_sections = set(old_data.keys()) | set(new_data.keys())
    for section in sorted(all_sections):
        if section in known_sections:
            continue
        old_sig = _section_signature(old_data.get(section) or [])
        new_sig = _section_signature(new_data.get(section) or [])
        added = [name for name in new_sig if name not in old_sig]
        removed = [name for name in old_sig if name not in new_sig]
        changed = [name for name in new_sig if name in old_sig and new_sig[name] != old_sig[name]]
        if added or removed or changed:
            other_changes.append({
                "section": section,
                "added": added,
                "removed": removed,
                "changed": changed,
            })

    return {
        "summary": {
            "added_policies": len(added_policies),
            "removed_policies": len(removed_policies),
            "changed_policies": len(changed_policies),
            "added_objects": len(added_objects),
            "removed_objects": len(removed_objects),
            "other_changes": len(other_changes),
        },
        "added_policies": added_policies,
        "removed_policies": removed_policies,
        "changed_policies": changed_policies,
        "added_objects": added_objects,
        "removed_objects": removed_objects,
        "other_changes": other_changes,
    }
import json
from pathlib import Path
from typing import Any

from flask import Flask, jsonify, render_template, request, send_from_directory, send_file

from app.parsers.config_parser import FortiGateConfigParser
from app.parsers.policy_csv_parser import PolicyStatsCsvParser
from app.parsers.runtime_parser import RuntimeStatsParser
from app.services.policy_renderer import build_view_model

from app.services.workbook_exporter import build_workbook
from datetime import date as _date
from app.services.severity_engine import evaluate_severity
from app.services.ai_analyzer import analyze_policies, analyze_single_policy, check_ollama_available
from app.services.license_checker import activate, is_licensed, get_license_info
from io import BytesIO

import sys as _sys
APO_VERSION = "v92-2026-08-11"
if getattr(_sys, 'frozen', False) and hasattr(_sys, '_MEIPASS'):
    BASE_DIR = Path(_sys._MEIPASS)
else:
    BASE_DIR = Path(__file__).resolve().parent.parent


def _stats_count(runtime_stats) -> int:
    """사용량 통계 건수. 평평한 형태와 종류별 형태를 모두 센다."""
    if not isinstance(runtime_stats, dict):
        return 0
    if "firewall" in runtime_stats or "proxy" in runtime_stats:
        # 섞인 dict에서 평평한 항목의 '필드 수'를 정책 수로 세지 않도록 한정한다.
        return sum(len(runtime_stats.get(k) or {}) for k in ("firewall", "proxy"))
    return len(runtime_stats)


def _config_sha(raw: str) -> str:
    """설정 원문의 SHA-256 앞 16자. 장비/스냅샷을 가리키는 안정된 식별자다.

    hostname은 공장 기본값 'FortiGate' 그대로인 장비가 흔해 서로 다른 고객
    장비가 같은 이름을 갖는다. 화면이 장비별로 기억하는 것(작업 완료 체크 등)을
    hostname에 걸면 그 장비들끼리 상태가 섞인다.
    """
    import hashlib
    return hashlib.sha256((raw or "").encode("utf-8", "replace")).hexdigest()[:16]


print(f"[APO] Version: {APO_VERSION}", flush=True)
print(f"[APO] Frozen: {getattr(_sys, 'frozen', False)}", flush=True)
print(f"[APO] BASE_DIR: {BASE_DIR}", flush=True)
print(f"[APO] Templates: {BASE_DIR / 'app' / 'templates'} "
      f"(exists={(BASE_DIR / 'app' / 'templates' / 'index.html').exists()})", flush=True)

IMPORT_DIR = BASE_DIR / "imports"
EXPORT_DIR = BASE_DIR / "exports"
DATA_DIR = BASE_DIR / "data"

APP_PORT = 5000
MAX_UPLOAD_BYTES = 50 * 1024 * 1024   # 50 MB

# 로컬 전용 앱이므로 자기 자신 외의 Host/Origin은 받지 않는다.
_ALLOWED_HOSTS = frozenset(
    f"{h}:{APP_PORT}" for h in ("127.0.0.1", "localhost", "[::1]")
) | frozenset(("127.0.0.1", "localhost", "[::1]"))
_ALLOWED_ORIGINS = frozenset(
    f"{scheme}://{h}" for scheme in ("http", "https")
    for h in (f"127.0.0.1:{APP_PORT}", f"localhost:{APP_PORT}", f"[::1]:{APP_PORT}")
)


def _license_required():
    """라이선스 미보유 시 402 응답을 반환, 보유 시 None.
    유료 Export 계열 라우트의 서버측 게이트 (클라이언트 게이트 우회 차단)."""
    if not is_licensed():
        return jsonify({"error": "License required. Please activate your Export license."}), 402
    return None


def create_app() -> Flask:
    app = Flask(
        __name__,
        template_folder=str(BASE_DIR / "app" / "templates"),
        static_folder=str(BASE_DIR / "app" / "static"),
    )

    # 업로드 상한. 미설정 시 설정 파일 하나로 수 GB 메모리를 쓸 수 있다
    # (파싱 결과가 입력 대비 수십 배로 부풀고, 응답으로 한 번 더 직렬화된다).
    app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES

    @app.errorhandler(413)
    def _too_large(_e):
        limit_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
        return jsonify({"error": f"File too large. Limit is {limit_mb} MB."}), 413

    @app.before_request
    def _guard_origin():
        """DNS 리바인딩 / 크로스 오리진 요청 차단.

        APO는 127.0.0.1에 바인딩되지만 그것만으로는 브라우저를 통한 접근을
        막지 못한다. 공격자 페이지가 짧은 TTL로 자기 도메인을 127.0.0.1로
        재바인딩하면 브라우저는 그 페이지를 APO와 '같은 오리진'으로 취급해,
        저장된 FortiGate 관리자 토큰을 읽거나 실제 정책을 변경할 수 있다.

        - Host 화이트리스트: 재바인딩된 요청은 Host가 공격자 도메인이라 걸린다.
        - Origin 검사: multipart는 CORS 프리플라이트가 없어 JSON 라우트와 달리
          크로스 오리진 POST가 그대로 도달하므로 여기서 막는다.
        """
        host = (request.host or "").lower()
        if host not in _ALLOWED_HOSTS:
            return jsonify({"error": "Invalid Host header"}), 403

        origin = request.headers.get("Origin")
        if origin and origin.lower() not in _ALLOWED_ORIGINS:
            return jsonify({"error": "Cross-origin request rejected"}), 403
        return None

    from datetime import datetime as _dt

    @app.context_processor
    def inject_year():
        return {'current_year': _dt.now().year}

    # ── 판정 프로파일 ──────────────────────────────────────────────────────
    # 시작 시 한 번 로드한다(APO_PROFILE 환경변수 또는 default + 구형
    # customer_rules.json 병합). UI에서 넘어온 임계값 오버라이드는
    # thresholds_override에 보관해 classify와 export가 같은 기준을 쓰게 한다.
    from app.services.profile_loader import (
        load_profile, engine_context, describe_inactive_rules,
    )
    app.config['profile'] = load_profile()
    app.config['thresholds_override'] = {}

    # UI가 조정할 수 있는 임계값. 이 밖의 키는 무시한다(임의 키 주입 방지).
    _THRESHOLD_FIELDS = {
        "dormancy_days":              ("int",  1, 3650),
        "long_dormancy_days":         ("int",  1, 7300),
        "use_absolute_hit_threshold": ("bool", None, None),
        "su_hit_multiplier":          ("int",  1, 100000),
        "ss_hit_threshold":           ("int",  1, 1000000),
        "ss_schedule_age_years":      ("int",  1, 20),
        "registration_fallback_year": ("int",  2000, 2100),
    }

    def _sanitize_thresholds(raw: dict) -> dict:
        out = {}
        for key, (kind, lo, hi) in _THRESHOLD_FIELDS.items():
            if key not in raw:
                continue
            val = raw[key]
            if val is None:
                out[key] = None       # 명시적 null = 해당 판정 끔
                continue
            if kind == "bool":
                out[key] = bool(val)
                continue
            try:
                num = int(val)
            except (TypeError, ValueError):
                continue
            if lo <= num <= hi:
                out[key] = num
        return out

    def _engine_ctx() -> dict:
        ctx = engine_context(app.config['profile'])
        override = app.config.get('thresholds_override') or {}
        if override:
            ctx["thresholds"] = {**ctx["thresholds"], **override}
        return ctx

    @app.get("/api/profile")
    def get_profile():
        profile = app.config['profile']
        merged = _engine_ctx()
        return jsonify({
            "name": (profile.get("meta") or {}).get("name", "default"),
            "description": (profile.get("meta") or {}).get("description", ""),
            "thresholds": merged["thresholds"],
            "rules": merged["rules"],
            "inactive_rules": describe_inactive_rules(profile),
        })

    for d in (IMPORT_DIR, EXPORT_DIR, DATA_DIR):
        d.mkdir(parents=True, exist_ok=True)

    @app.get("/")
    def index():
        return render_template("index.html")

    @app.get("/version")
    def version():
        # 절대경로 등 내부 파일시스템 정보는 노출하지 않는다.
        return jsonify({
            "version": APO_VERSION,
            "frozen": bool(getattr(_sys, 'frozen', False)),
        })

    @app.post("/api/config/parse")
    def parse_config():
        uploaded = request.files.get("config_file")
        if not uploaded:
            return jsonify({"error": "config_file is required"}), 400

        from werkzeug.utils import secure_filename
        filename = secure_filename(uploaded.filename or "") or "fortigate.conf"
        save_path = IMPORT_DIR / filename
        uploaded.save(save_path)

        raw = save_path.read_text(encoding="utf-8", errors="ignore")
        parser = FortiGateConfigParser(raw)
        parsed = parser.parse()
        view = build_view_model(parsed, runtime_stats={})

        export_name = f"{Path(filename).stem}.parsed.json"
        export_path = EXPORT_DIR / export_name
        export_path.write_text(json.dumps({"parsed": parsed, "view": view}, indent=2), encoding="utf-8")
        app.config['last_parsed'] = parsed
        app.config['last_raw_config'] = raw      # Version Advisor 기능 감지용
        app.config['last_config_filename'] = filename   # 리포트 문서정보용
        app.config['last_runtime_stats'] = {}   # 새 Config 로드 시 CSV stats 초기화
        # 새 config는 다른 장비일 수 있다. 이전 장비의 DNS 캐시가 남으면
        # 그 스냅샷으로 새 장비를 판정하는 혼선이 생긴다(재비판). 함께 비운다.
        app.config.pop('fqdn_cache', None)

        return jsonify(
            {
                "message": "Config parsed successfully",
                "filename": filename,
                # 이 설정을 가리키는 안정된 식별자. hostname은 공장 기본값
                # "FortiGate"로 남아 있는 장비가 흔해 장비 구분에 못 쓴다.
                "config_sha": _config_sha(raw),
                "parsed": parsed,
                "view": view,
                "export_json": export_name,
            }
        )

    @app.post("/api/runtime/import")
    def import_runtime():
        parser = RuntimeStatsParser()
        merged: dict[str, dict[str, Any]] = {}

        uploaded_files = request.files.getlist("runtime_files")
        pasted_text = request.form.get("runtime_text", "")

        for file in uploaded_files:
            text = file.read().decode("utf-8", errors="ignore")
            stats = parser.parse_text(text)
            merged.update(stats)

        if pasted_text.strip():
            stats = parser.parse_text(pasted_text)
            merged.update(stats)

        # 이 경로(CLI 출력 붙여넣기)는 정책 종류를 알 수 없다. 이미 종류별로
        # 나뉜 상태에 평평하게 합치면 _stats_for가 네임스페이스만 읽어 방금 넣은
        # 통계를 통째로 버린다 — 조용히 '사용량 없음'이 되어 미사용 판정으로
        # 이어지므로 가장 위험한 실패다. 종류를 모르면 양쪽에 넣는다.
        existing = app.config.get('last_runtime_stats') or {}
        if "firewall" in existing or "proxy" in existing:
            for k in ("firewall", "proxy"):
                existing.setdefault(k, {}).update(merged)
        else:
            existing.update(merged)
        app.config['last_runtime_stats'] = existing

        return jsonify({"runtime_stats": merged})

    @app.post("/api/policy-stats/import")
    def import_policy_stats_csv():
        parser = PolicyStatsCsvParser()
        merged: dict[str, dict[str, Any]] = {}
        uploaded_files = request.files.getlist("policy_stats_files")
        pasted_text = request.form.get("policy_stats_text", "")

        for file in uploaded_files:
            text = file.read().decode("utf-8", errors="ignore")
            stats = parser.parse_text(text)
            merged.update(stats)

        if pasted_text.strip():
            stats = parser.parse_text(pasted_text)
            merged.update(stats)

        # 정책 종류별로 담는다. 평평하게 합치면 FW CSV와 Proxy CSV의 같은 번호가
        # 서로를 덮어써, 어느 파일을 먼저 올렸느냐에 따라 판정이 달라졌다.
        # 종류를 안 알려주면(구형 클라이언트·직접 호출) 예전처럼 평평하게 둔다.
        ptype = str(request.form.get("policy_type", "")).strip().lower()
        ns = {"fw": "firewall", "firewall": "firewall",
              "proxy": "proxy", "proxy-policy": "proxy"}.get(ptype)
        existing = app.config.get('last_runtime_stats') or {}
        if ns:
            if not ("firewall" in existing or "proxy" in existing):
                # 평평한 옛 상태 → 종류별로 승격. 기존 값을 버리면 사용자가 앞서
                # 붙여넣은 통계가 조용히 사라진다. 평평했다는 건 "종류를 몰라
                # 양쪽에 적용"이라는 뜻이므로 그대로 양쪽에 복사한 뒤 덮어쓴다.
                existing = {"firewall": dict(existing), "proxy": dict(existing)}
            existing.setdefault(ns, {}).update(merged)
        else:
            if "firewall" in existing or "proxy" in existing:
                # 이미 종류별인데 종류 없는 입력이 오면 양쪽에 넣는 수밖에 없다
                for k in ("firewall", "proxy"):
                    existing.setdefault(k, {}).update(merged)
            else:
                existing.update(merged)
        app.config['last_runtime_stats'] = existing
        summary = {
            "count": len(merged),
            # int/str 혼합 정렬 금지 — CSV에 비숫자 ID가 한 건이라도 있으면
            # int(x)와 str이 비교돼 TypeError 500이 났다(감사 C1). 튜플 키로 분리.
            "matched_policy_ids": sorted(
                merged.keys(),
                key=lambda x: (0, int(x)) if str(x).isdigit() else (1, str(x))),
        }
        return jsonify({"runtime_stats": merged, "summary": summary})

    @app.post("/api/policies/render")
    def render_policies():
        payload = request.get_json(silent=True) or {}
        parsed = payload.get("parsed")
        runtime_stats = payload.get("runtime_stats", {})
        if not parsed:
            return jsonify({"error": "parsed payload is required"}), 400

        view = build_view_model(parsed, runtime_stats)
        return jsonify({"view": view})

    @app.post("/api/export/workbook")
    def export_workbook():
        gate = _license_required()
        if gate:
            return gate
        from werkzeug.utils import secure_filename
        payload = request.get_json(silent=True) or {}
        sheets = payload.get("sheets") or {}
        # secure_filename은 str만 받는다. JSON으로 null/숫자/객체가 오면
        # TypeError가 try 밖에서 터져 500이 되고, 콘솔에 절대경로 트레이스백이 남는다.
        raw_name = payload.get("workbook_name") or "firewall_policy_optimizer_export"
        workbook_name = secure_filename(str(raw_name)) or "firewall_policy_optimizer_export"
        output_path = EXPORT_DIR / f"{workbook_name}.xlsx"
        try:
            build_workbook(output_path, sheets)
        except Exception as exc:
            print(f"[APO] workbook export failed: {exc}", flush=True)
            return jsonify({"error": "Failed to build workbook export"}), 400
        return send_file(output_path, as_attachment=True, download_name=output_path.name)

    @app.get("/exports/<path:filename>")
    def download_export(filename: str):
        from werkzeug.utils import secure_filename
        safe_name = secure_filename(filename)
        if not safe_name or safe_name != filename:
            return jsonify({"error": "Invalid filename"}), 400
        # 무료 산출물(파싱 결과 JSON)만 이 경로로 제공. 유료 산출물(.xlsx 등)은
        # 라이선스 게이트가 걸린 전용 라우트로만 내려가야 하므로 여기서 차단.
        if not safe_name.endswith(".parsed.json"):
            return jsonify({"error": "Not available"}), 404
        return send_from_directory(EXPORT_DIR, safe_name, as_attachment=True)

    @app.post("/api/config/diff")
    def api_config_diff():
        old_file = request.files.get("old_config")
        new_file = request.files.get("new_config")

        if not old_file or not new_file:
            return jsonify({"error": "Both baseline and target config files are required."}), 400

        try:
            old_text = old_file.read().decode("utf-8", errors="ignore")
            new_text = new_file.read().decode("utf-8", errors="ignore")

            old_data = parse_config_text(old_text)
            new_data = parse_config_text(new_text)

            result = _compute_config_diff(old_data, new_data)
            return jsonify(result)
        except Exception as exc:
            print(f"[APO] config diff failed: {exc}", flush=True)
            return jsonify({"error": "Failed to compare configs. Check that both files are valid FortiGate configs."}), 500

    app.config.setdefault('user_ranges', [])

    @app.get("/api/user-ranges")
    def get_user_ranges():
        return jsonify({"user_ranges": app.config.get('user_ranges', [])})

    @app.post("/api/user-ranges/set")
    def set_user_ranges():
        payload = request.get_json(silent=True) or {}
        ranges = [{"cidr": str(r).strip()} for r in payload.get("ranges", []) if str(r).strip()]
        app.config['user_ranges'] = ranges
        return jsonify({"ok": True, "count": len(ranges)})

    @app.post("/api/severity/classify")
    def severity_classify():
        # 임계값 오버라이드는 config 로드 여부와 무관하게 먼저 반영한다 —
        # 사용자가 설정을 조정해 두고 나중에 config를 올리는 순서도 유효하다.
        payload = request.get_json(silent=True) or {}
        if "thresholds" in payload:
            app.config['thresholds_override'] = _sanitize_thresholds(
                payload.get("thresholds") or {})
        parsed = app.config.get('last_parsed')
        if not parsed:
            return jsonify({"error": "No config loaded. Upload a config file first."}), 400
        service_groups = parsed.get("service_groups", {})
        user_ranges = app.config.get('user_ranges', [])
        context = {
            "service_groups": service_groups,
            "user_ranges": user_ranges,
            "today": _date.today(),
            **_engine_ctx(),
        }
        def classify_list(policies):
            result = []
            for p in (policies or []):
                sev = evaluate_severity(p, context)
                result.append({**p, **sev})
            return result

        from app.services.policy_renderer import build_view_model
        from app.services.reachability import detect_unreachable
        runtime_stats = app.config.get('last_runtime_stats', {})
        view = build_view_model(parsed, runtime_stats)
        return jsonify({
            "firewall": classify_list(view.get("firewall_policy", [])),
            "proxy":    classify_list(view.get("firewall_proxy_policy", [])),
            # 이번 분석에서 적용되지 않은 판정 규칙. export 시 Notes 시트로 실린다.
            "inactive_rules": describe_inactive_rules(app.config['profile']),
            # 정적 도달불가 검출. severity와 별개 축 — 등급을 바꾸지 않는다.
            "reachability": {
                **detect_unreachable(
                    parsed,
                    (app.config.get('fqdn_cache') or {}).get("map")),
                "fqdn_captured_at":
                    (app.config.get('fqdn_cache') or {}).get("captured_at"),
            },
        })

    # ── AI 분석 (Ollama/hermes3 로컬) ──────────────────────────────────────
    @app.get("/api/ai/status")
    def ai_status():
        available = check_ollama_available()
        return jsonify({"available": available, "model": "hermes3:latest"})

    @app.post("/api/ai/analyze")
    def ai_analyze():
        payload = request.get_json(silent=True) or {}
        policies = payload.get("policies", [])
        mode = payload.get("mode", "summary")   # "summary" | "single"
        model = payload.get("model", "hermes3:latest")
        try:
            if mode == "single" and policies:
                result = analyze_single_policy(policies[0], model)
            else:
                result = analyze_policies(policies, model)
            return jsonify({"result": result, "model": model})
        except RuntimeError as e:
            return jsonify({"error": str(e)}), 503

    @app.post("/api/export/severity-workbook")
    def export_severity_workbook():
        gate = _license_required()
        if gate:
            return gate
        # config 없이 내려가는 '빈 감사 보고서'는 문서정보(SHA·장비)가 공란인
        # 채 형식만 갖춰서 더 위험하다(감사 C2). 다른 분석 라우트와 동일하게 막는다.
        if not app.config.get('last_parsed'):
            return jsonify({"error": "No config loaded. Upload a config file first."}), 400
        from app.services.workbook_exporter import build_severity_workbook
        payload = request.get_json(silent=True) or {}
        # 문서 정보는 서버가 주입한다 — 감사 보고서는 "무엇을, 어떤 기준으로,
        # 언제 분석했나"를 스스로 증명해야 하고, 그 사실은 클라이언트가 아니라
        # 서버가 안다. config SHA-256이 있어야 "이 보고서는 그 설정 파일에
        # 대한 것"이라는 대응이 성립한다.
        import hashlib as _hashlib
        raw_cfg = app.config.get('last_raw_config') or ''
        meta = (app.config.get('last_parsed') or {}).get('meta', {})
        merged_ctx = _engine_ctx()
        payload["report_meta"] = {
            "apo_version": APO_VERSION,
            "generated_at": _dt.now().strftime("%Y-%m-%d %H:%M"),
            "hostname": meta.get("hostname") or "",
            "config_version": meta.get("config_version") or "",
            "buildno": meta.get("buildno") or "",
            "config_filename": app.config.get('last_config_filename') or "",
            "config_sha256": _hashlib.sha256(raw_cfg.encode("utf-8", "ignore")).hexdigest() if raw_cfg else "",
            "profile": (app.config['profile'].get("meta") or {}).get("name", "default"),
            "thresholds": merged_ctx["thresholds"],
            "rules": merged_ctx["rules"],
            "user_range_count": len(app.config.get('user_ranges', [])),
            # 종류별 dict는 비어 있어도 truthy다({"firewall":{},"proxy":{}}).
            # 실제 건수를 봐야 "사용량 데이터 있음"이 사실이 된다.
            "csv_loaded": _stats_count(app.config.get('last_runtime_stats')) > 0,
        }
        if _branding_allowed():
            b = _load_branding()
            if b.get("company"):
                payload["report_meta"]["branding"] = b
        try:
            xlsx_bytes = build_severity_workbook(payload)
        except Exception as exc:
            print(f"[APO] severity workbook export failed: {exc}", flush=True)
            return jsonify({"error": "Failed to build severity export"}), 400
        return app.response_class(
            response=xlsx_bytes,
            status=200,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={"Content-Disposition": "attachment; filename=severity_export.xlsx"}
        )

    # ── 샘플 리포트 (무료) ────────────────────────────────────────────────
    # 이 도구의 가치는 결국 엑셀 산출물인데, 지금까지는 **사기 전에 그걸 볼
    # 방법이 없었다**. 번들된 데모 설정으로 같은 파이프라인을 돌려 실제 워크북을
    # 그대로 내려준다. 라이선스 게이트를 걸지 않는다 — 게이트를 걸면 "사기 전엔
    # 못 본다"는 문제가 그대로다.
    @app.get("/api/sample-report")
    def sample_report():
        """번들 데모 설정으로 만든 실제 감사 워크북.

        현재 세션 상태(last_parsed 등)를 **절대 건드리지 않는다**. 사용자가
        고객 설정을 열어 둔 채 샘플을 눌렀다가 작업이 날아가면 안 된다.
        """
        from app.services.workbook_exporter import build_severity_workbook
        from app.services.policy_renderer import build_view_model
        from app.services.reachability import detect_unreachable
        from app.parsers.policy_csv_parser import PolicyStatsCsvParser
        import hashlib as _hashlib

        # DATA_DIR은 사용자 데이터 경로다. 번들 자산은 app/data에 있고, frozen
        # 상태에서는 _MEIPASS 아래로 풀린다(version_advisor와 같은 규칙).
        if getattr(_sys, 'frozen', False) and hasattr(_sys, '_MEIPASS'):
            bundled = Path(_sys._MEIPASS) / "app" / "data"
        else:
            bundled = Path(__file__).resolve().parent / "data"
        cfg_path = bundled / "sample_config.conf"
        csv_path = bundled / "sample_policy_stats.csv"
        if not cfg_path.exists():
            return jsonify({"error": "Sample configuration is not bundled in this build."}), 404
        raw = cfg_path.read_text(encoding="utf-8", errors="ignore")
        parsed = FortiGateConfigParser(raw).parse()

        # 샘플 CSV는 firewall 정책 통계다. 평평하게 넘기면 proxy 정책이 같은
        # 번호의 firewall 통계를 빌려가, v91이 고친 교차오염을 **샘플 산출물에서
        # 그대로 재현**한다(샘플 config에 proxy-policy 1·2가 있다).
        runtime_stats = {"firewall": {}, "proxy": {}}
        if csv_path.exists():
            try:
                runtime_stats["firewall"] = PolicyStatsCsvParser().parse_text(
                    csv_path.read_text(encoding="utf-8-sig", errors="ignore"))
            except Exception:
                runtime_stats["firewall"] = {}

        # 샘플이 '판정 불가'로 도배되지 않도록 데모 설정의 사용자 대역을 준다.
        context = {
            "service_groups": parsed.get("service_groups", {}),
            "user_ranges": [{"cidr": "10.10.0.0/16"}],
            "today": _date.today(),
            **_engine_ctx(),
        }
        view = build_view_model(parsed, runtime_stats)
        classify = lambda ps: [{**p, **evaluate_severity(p, context)} for p in (ps or [])]
        meta = parsed.get("meta", {})
        payload = {
            "firewall": classify(view.get("firewall_policy", [])),
            "proxy": classify(view.get("firewall_proxy_policy", [])),
            "inactive_rules": describe_inactive_rules(app.config['profile']),
            "reachability": {**detect_unreachable(parsed, None),
                             "fqdn_captured_at": None},
            "report_meta": {
                "apo_version": APO_VERSION,
                "generated_at": _dt.now().strftime("%Y-%m-%d %H:%M"),
                "hostname": meta.get("hostname") or "DEMO-FGT-01",
                "config_version": meta.get("config_version") or "",
                "buildno": meta.get("buildno") or "",
                "config_filename": "sample_config.conf",
                "config_sha256": _hashlib.sha256(raw.encode("utf-8", "ignore")).hexdigest(),
                "profile": (app.config['profile'].get("meta") or {}).get("name", "default"),
                "thresholds": context["thresholds"],
                "rules": context["rules"],
                "user_range_count": 1,
                # {"firewall":{},"proxy":{}}는 비어 있어도 truthy다. 번들 CSV가
                # 빠지거나 파싱에 실패해도 "사용량 있음"이라고 적히면 안 된다.
                "csv_loaded": _stats_count(runtime_stats) > 0,
                # 이게 실제 장비 보고서로 오해되면 안 된다.
                "sample": True,
            },
        }
        try:
            xlsx_bytes = build_severity_workbook(payload)
        except Exception as exc:
            print(f"[APO] sample report failed: {exc}", flush=True)
            return jsonify({"error": "Failed to build the sample report"}), 500
        return app.response_class(
            response=xlsx_bytes, status=200,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={"Content-Disposition": "attachment; filename=APO_sample_report.xlsx"})

    # ── 정확도 검증 워크플로우 ─────────────────────────────────────────────
    # 의도적으로 무료다: 이 기능의 목적은 "APO 판정을 그대로 믿지 말고
    # 표본으로 직접 검증하라"는 신뢰 구축이고, 구매 전 사용자가 해볼 수
    # 있어야 의미가 있다. 감사 산출물(유료 export)과는 성격이 다르다.
    @app.post("/api/verification/sample")
    def verification_sample():
        parsed = app.config.get('last_parsed')
        if not parsed:
            return jsonify({"error": "No config loaded. Upload a config file first."}), 400
        from app.services import verification as vf
        from app.services.policy_renderer import build_view_model as _bvm
        fmt = str(request.args.get("format", "xlsx")).lower()
        context = {
            "service_groups": parsed.get("service_groups", {}),
            "user_ranges": app.config.get('user_ranges', []),
            "today": _date.today(),
            **_engine_ctx(),
        }
        view = _bvm(parsed, app.config.get('last_runtime_stats', {}))
        results = []
        for key, ptype in (("firewall_policy", "firewall"),
                           ("firewall_proxy_policy", "proxy")):
            for pol in view.get(key, []):
                results.append({**pol, "policy_type": ptype,
                                **evaluate_severity(pol, context)})
        samples = vf.stratified_sample(results)
        rows = vf.build_rows(samples)
        if fmt == "csv":
            return app.response_class(
                response=vf.to_csv(rows), status=200, mimetype="text/csv",
                headers={"Content-Disposition":
                         "attachment; filename=apo_verification_sample.csv"})
        return app.response_class(
            response=vf.to_xlsx(rows), status=200,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition":
                     "attachment; filename=apo_verification_sample.xlsx"})

    @app.post("/api/verification/score")
    def verification_score():
        from app.services import verification as vf
        uploaded = request.files.get("worksheet")
        if not uploaded:
            return jsonify({"error": "worksheet file is required (.csv or .xlsx)"}), 400
        labels = {label: key for key, label in vf.WORKSHEET_COLUMNS}
        rows = []
        name = (uploaded.filename or "").lower()
        try:
            if name.endswith(".xlsx"):
                import io as _io
                from openpyxl import load_workbook
                wb = load_workbook(_io.BytesIO(uploaded.read()), read_only=True)
                ws = wb.active
                header = None
                for raw in ws.iter_rows(values_only=True):
                    if header is None:
                        header = [labels.get(str(h or "").strip(), str(h or "").strip())
                                  for h in raw]
                        continue
                    rows.append({header[i]: ("" if v is None else str(v))
                                 for i, v in enumerate(raw) if i < len(header)})
            else:
                import csv as _csv, io as _io
                raw_bytes = uploaded.read()
                # 한국 Excel의 "CSV(쉼표로 분리)" 기본 저장은 cp949다. 강제
                # utf-8(errors=replace)로 읽으면 한글 판정값("일치")이 U+FFFD로
                # 깨져 전량 '미기입'이 되는데, 오류 없이 0건 채점으로 위장된다
                # (재비판 NEW-BUG). 엄격 디코딩 + cp949 폴백으로 바꾼다.
                text = None
                for enc in ("utf-8-sig", "cp949"):
                    try:
                        text = raw_bytes.decode(enc)
                        break
                    except UnicodeDecodeError:
                        continue
                if text is None:
                    return jsonify({"error": "Could not decode the CSV. "
                                    "Save it as UTF-8 or the Korean Excel default (CP949)."}), 400
                first = text.splitlines()[0] if text.splitlines() else ""
                delim = ";" if first.count(";") > first.count(",") else ","
                for raw in _csv.DictReader(_io.StringIO(text), delimiter=delim):
                    rows.append({labels.get((k or "").strip(), (k or "").strip()): v
                                 for k, v in raw.items()})
        except Exception:
            return jsonify({"error": "Could not read the worksheet. "
                            "Upload the file generated by APO (.csv or .xlsx)."}), 400
        if not rows:
            return jsonify({"error": "The worksheet has no data rows."}), 400
        # 헤더가 판정 컬럼으로 매핑되지 않았으면(인코딩·구분자·양식 문제) 빈
        # 정확도 리포트로 위장하지 말고 원인을 말한다 (재비판 SUSPECT 2건).
        if not any("verdict" in r for r in rows):
            return jsonify({"error": "Worksheet columns were not recognized — "
                            "upload the sheet generated by APO, keeping its header row."}), 400
        stats = vf.compute_precision(rows)
        return jsonify(stats)

    # ── 원클릭 장비 수집 (v83) — REST API로 config·통계·FQDN 일괄 ─────────
    # 파일 업로드의 대안(선택). 등록된 장비(remediation과 같은 연결 정보)에서
    # 읽기 전용으로 수집해 기존 분석 파이프라인에 그대로 주입한다.
    def _resolve_device(payload: dict):
        """요청의 연결 정보 또는 등록된 장비를 반환. (device, error_response)

        두 라우트가 같은 규칙을 쓰도록 한 곳에 둔다 — 같은 로직을 복사하면
        한쪽만 고쳐지는 사고가 난다(오늘 캐시 무효화에서 겪은 유형).
        """
        device = app.config.get('remediation_device')
        if payload.get("ip"):
            from app.services.remediation_service import validate_device_address
            try:
                port = int(payload.get("port", 443) or 443)
                if not (1 <= port <= 65535):
                    raise ValueError("port out of range")
                device = {
                    "ip": validate_device_address(payload.get("ip", "")),
                    "port": port,
                    "token": str(payload.get("token", "")).strip(),
                    "vdom": str(payload.get("vdom", "root")).strip() or "root",
                    "verify_ssl": bool(payload.get("verify_ssl", True)),
                }
            except ValueError as exc:
                return None, (jsonify({"error": str(exc) or
                                       "Invalid device address or port"}), 400)
        if not device or not device.get("token"):
            return None, (jsonify({
                "error": "No device is connected. Register one in the "
                         "Remediation tab, or use 'Collect from device' with "
                         "an IP and API token."}), 400)
        return device, None

    @app.post("/api/collect/device")
    def collect_device():
        payload = request.get_json(silent=True) or {}
        device, err = _resolve_device(payload)
        if err:
            return err

        from app.services.device_collector import collect_from_device
        import requests as _req
        try:
            data = collect_from_device(device)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except _req.exceptions.SSLError:
            return jsonify({"error": "TLS verification failed. Fix the device "
                            "certificate, or disable verification only on a "
                            "trusted management network."}), 502
        except _req.exceptions.RequestException:
            return jsonify({"error": "Could not reach the device. Check the IP, "
                            "port, and that REST API access is enabled."}), 502

        # 수집물을 파일 업로드와 동일한 상태 슬롯에 넣는다 — 이후 classify·
        # export가 출처(파일 vs 장비)를 구분할 필요 없이 그대로 동작한다.
        parsed = FortiGateConfigParser(data["config_text"]).parse()
        app.config['last_parsed'] = parsed
        app.config['last_raw_config'] = data["config_text"]
        app.config['last_config_filename'] =             (parsed.get("meta") or {}).get("hostname", "device") + ".conf"
        app.config['last_runtime_stats'] = data["runtime_stats"]
        # 새 장비를 수집하면 이전 장비의 DNS 스냅샷은 무효다. 업로드 경로와
        # 같은 불변식 — 수집이 FQDN을 못 가져왔다고 이전 캐시를 남기면
        # 장비 A의 해석으로 장비 B를 판정하는 오탐이 난다(재비판 NEW-BUG).
        app.config.pop('fqdn_cache', None)
        if data["fqdn_map"]:
            app.config['fqdn_cache'] = {
                "map": data["fqdn_map"],
                "captured_at": _dt.now().strftime("%Y-%m-%d %H:%M"),
            }
        from app.services.policy_renderer import build_view_model as _bvm
        view = _bvm(parsed, data["runtime_stats"])
        warnings = list(data["warnings"])
        if not device.get("verify_ssl", True):
            warnings.insert(0, "TLS verification was disabled for this "
                            "collection — the API token traveled over an "
                            "unverified channel")
        return jsonify({
            "ok": True,
            "hostname": (parsed.get("meta") or {}).get("hostname", ""),
            "config_sha": _config_sha(data["config_text"]),
            "policies": len(view.get("firewall_policy", [])),
            # 화면은 이 값으로 "사용량 데이터가 있는가"를 판단한다. 개수만 주고
            # 본문을 빼면, 장비에서 통계를 잘 받아온 직후에도 화면은 "사용량
            # 없음"이라고 경고한다(재비판 3회전).
            "runtime_stats": data["runtime_stats"],
            # 종류별로 나뉘어 있으므로 합계를 센다(예전엔 평평한 dict였다).
            "stats_count": _stats_count(data["runtime_stats"]),
            "fqdn_count": len(data["fqdn_map"]),
            "warnings": warnings,
            "view": view,
            "parsed": parsed,
        })

    # ── FQDN 캐시 (B-10) — dnsproxy 덤프로 FQDN 정책 판정 포함 ────────────
    @app.post("/api/fqdn-cache/import")
    def fqdn_cache_import():
        from app.parsers.fqdn_cache_parser import parse_dnsproxy_dump
        uploaded = request.files.get("dump")
        if not uploaded:
            return jsonify({"error": "dump file is required — output of "
                            "'diagnose test application dnsproxy 6'"}), 400
        text = uploaded.read().decode("utf-8", errors="ignore")
        mapping = parse_dnsproxy_dump(text)
        if not mapping:
            # 왜 실패했는지 알 수 있게 진단을 함께 준다. "인식 못 했다"만
            # 던지면 사용자가 파일이 잘못된 건지 도구가 못 읽는 건지 모른다.
            lines = text.splitlines()
            import re as _re
            has_dom = any(_re.search(r'[A-Za-z0-9\-]+\.[A-Za-z]{2,}', ln)
                          for ln in lines[:2000])
            has_ip = any(_re.search(r'\d{1,3}(?:\.\d{1,3}){3}', ln)
                         for ln in lines[:2000])
            hint = ("The file has domains and IPs but not in a layout this "
                    "parser recognizes — please send a few sample lines so "
                    "it can be supported."
                    if (has_dom and has_ip) else
                    "The file does not appear to contain FQDN entries. Run "
                    "'diagnose test application dnsproxy 6' (or 'diagnose "
                    "firewall fqdn list') and save the full console output.")
            return jsonify({
                "error": f"No FQDN→IP entries recognized. {hint}",
                "diagnostics": {"lines": len(lines),
                                "domains_seen": has_dom,
                                "ipv4_seen": has_ip},
            }), 400
        app.config['fqdn_cache'] = {
            "map": mapping,
            "captured_at": _dt.now().strftime("%Y-%m-%d %H:%M"),
        }
        return jsonify({"ok": True, "names": len(mapping),
                        "ips": sum(len(v) for v in mapping.values()),
                        "captured_at": app.config['fqdn_cache']["captured_at"]})

    @app.post("/api/fqdn-cache/fetch-device")
    def fqdn_cache_fetch_device():
        """연결된 장비에서 FQDN 해석만 가져온다 — 덤프 파일 없이."""
        payload = request.get_json(silent=True) or {}
        device, err = _resolve_device(payload)
        if err:
            return err
        from app.services.device_collector import fetch_fqdn_map
        import requests as _req
        try:
            mapping = fetch_fqdn_map(device)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        except _req.exceptions.SSLError:
            return jsonify({"error": "TLS verification failed."}), 502
        except _req.exceptions.RequestException:
            return jsonify({"error": "Could not reach the device."}), 502
        if not mapping:
            return jsonify({"error": "The device reported no resolved FQDN "
                            "addresses. FQDN objects resolve only after they "
                            "have been used."}), 400
        app.config['fqdn_cache'] = {
            "map": mapping,
            "captured_at": _dt.now().strftime("%Y-%m-%d %H:%M"),
        }
        return jsonify({"ok": True, "names": len(mapping),
                        "ips": sum(len(v) for v in mapping.values()),
                        "captured_at": app.config['fqdn_cache']["captured_at"]})

    # ── 감사 증적 (B-06) — 스냅샷 시계열 → Review Evidence ────────────────
    @app.post("/api/audit/timeline")
    def audit_timeline():
        from app.services import audit_evidence as ae
        files = request.files.getlist("snapshots")
        if len(files) < 2:
            return jsonify({"error": "Upload at least two config snapshots, "
                            "oldest first."}), 400
        if len(files) > 24:
            return jsonify({"error": "At most 24 snapshots per timeline."}), 400
        snaps, parsed_list = [], []
        for f in files:
            raw = f.read()
            try:
                parsed = FortiGateConfigParser(
                    raw.decode("utf-8", errors="ignore")).parse()
            except Exception:
                return jsonify({"error": f"Could not parse snapshot "
                                f"'{f.filename}'."}), 400
            snaps.append(ae.snapshot_meta(f.filename or "", raw, parsed))
            parsed_list.append(parsed)
        intervals = []
        for i in range(len(parsed_list) - 1):
            diff = _compute_config_diff(parsed_list[i], parsed_list[i + 1])
            intervals.append({
                "from": snaps[i]["filename"],
                "to": snaps[i + 1]["filename"],
                "summary": ae.interval_summary(diff),
            })
        timeline = {"snapshots": snaps, "intervals": intervals}
        # 증적 export가 같은 데이터를 쓰도록 보관 (재업로드 불필요)
        app.config['last_audit_timeline'] = timeline
        return jsonify(timeline)

    @app.post("/api/audit/evidence-workbook")
    def audit_evidence_workbook():
        gate = _license_required()   # 감사 산출물 = 유료 export 계열
        if gate:
            return gate
        timeline = app.config.get('last_audit_timeline')
        if not timeline:
            return jsonify({"error": "Run the snapshot timeline first."}), 400
        from app.services import audit_evidence as ae
        rm = {"apo_version": APO_VERSION,
              "generated_at": _dt.now().strftime("%Y-%m-%d %H:%M")}
        if _branding_allowed():
            b = _load_branding()
            if b.get("company"):
                rm["branding"] = b
        try:
            xlsx = ae.build_evidence_workbook(timeline, rm)
        except Exception as exc:
            print(f"[APO] evidence workbook failed: {exc}", flush=True)
            return jsonify({"error": "Failed to build evidence workbook"}), 400
        return app.response_class(
            response=xlsx, status=200,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition":
                     "attachment; filename=apo_review_evidence.xlsx"})

    # ── Version Advisor ────────────────────────────────────────────────────
    @app.get("/api/version-advisor")
    def version_advisor():
        from app.services.version_advisor import run_advisor
        parsed = app.config.get('last_parsed')
        if not parsed:
            return jsonify({"error": "No config loaded. Upload a config file first."}), 400
        meta = parsed.get("meta", {})
        raw_config = app.config.get('last_raw_config', '')
        result = run_advisor(meta, raw_config)
        return jsonify(result)

    # ── Remediation ───────────────────────────────────────────────────────
    app.config.setdefault('remediation_device', {})

    @app.post("/api/remediation/device")
    def remediation_device_save():
        payload = request.get_json(silent=True) or {}
        try:
            port = int(payload.get("port", 443) or 443)
        except (TypeError, ValueError):
            return jsonify({"error": "port must be a number"}), 400
        if not (1 <= port <= 65535):
            return jsonify({"error": "port out of range"}), 400
        from app.services.remediation_service import validate_device_address
        try:
            device_ip = validate_device_address(payload.get("ip", ""))
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400

        device = {
            "ip":         device_ip,
            "port":       port,
            "token":      str(payload.get("token", "")).strip(),
            "vdom":       str(payload.get("vdom", "root")).strip() or "root",
            # 기본값 True — 명시적으로 끄지 않는 한 TLS 검증을 수행한다.
            "verify_ssl": bool(payload.get("verify_ssl", True)),
        }
        if not device["token"]:
            return jsonify({"error": "token is required"}), 400
        app.config['remediation_device'] = device
        return jsonify({"ok": True})

    @app.get("/api/remediation/device")
    def remediation_device_get():
        """등록된 장비를 알려준다 — 토큰은 절대 내보내지 않는다.

        되돌릴 수 없는 변경을 확인시키는 창이 입력칸 값을 읽으면, 사용자가 저장
        뒤에 칸만 고쳐 놓은 경우 실제 적용 대상과 다른 장비를 보여준다. 적용은
        서버가 들고 있는 이 값으로 나가므로, 확인창도 같은 값을 보여야 한다.
        """
        d = app.config.get('remediation_device') or {}
        return jsonify({
            "registered": bool(d.get("ip")),
            "ip": d.get("ip", ""),
            "port": d.get("port", 443),
            "vdom": d.get("vdom", ""),
        })

    @app.get("/api/remediation/device/test")
    def remediation_device_test():
        from app.services.remediation_service import test_connection
        device = app.config.get('remediation_device', {})
        if not device.get("ip"):
            return jsonify({"ok": False, "message": "No device registered"}), 400
        return jsonify(test_connection(device))

    @app.get("/api/remediation/candidates")
    def remediation_candidates():
        from app.services.remediation_service import get_candidates
        from app.services.severity_engine import evaluate_severity
        from app.services.policy_renderer import build_view_model
        parsed = app.config.get('last_parsed')
        if not parsed:
            return jsonify({"error": "No config loaded. Upload a config file first."}), 400
        service_groups = parsed.get("service_groups", {})
        user_ranges = app.config.get('user_ranges', [])
        context = {
            "service_groups": service_groups,
            "user_ranges": user_ranges,
            "today": _date.today(),
            **_engine_ctx(),
        }
        runtime_stats = app.config.get('last_runtime_stats', {})
        view = build_view_model(parsed, runtime_stats)
        def classify(policies):
            return [{**p, **evaluate_severity(p, context)} for p in (policies or [])]
        fw_classified  = classify(view.get("firewall_policy", []))
        prx_classified = classify(view.get("firewall_proxy_policy", []))
        severity_results = {"firewall": fw_classified, "proxy": prx_classified}

        # Severity 탭과 동일한 urgency 분포 계산 → 카운트 불일치 디버그용
        all_classified = fw_classified + prx_classified
        urgency_dist = {}
        for p in all_classified:
            u = p.get("urgency", 0)
            urgency_dist[str(u)] = urgency_dist.get(str(u), 0) + 1

        candidates = get_candidates(parsed, severity_results)
        candidates["_debug"] = {
            "total_policies": len(all_classified),
            "urgency_distribution": urgency_dist,
            "critical_high_total": urgency_dist.get("1", 0) + urgency_dist.get("2", 0),
        }
        return jsonify(candidates)

    @app.post("/api/remediation/apply")
    def remediation_apply():
        from app.services.remediation_service import disable_policies
        device = app.config.get('remediation_device', {})
        if not device.get("ip"):
            return jsonify({"error": "No device registered"}), 400
        payload = request.get_json(silent=True) or {}
        policies = payload.get("policies", [])
        if not policies:
            return jsonify({"error": "No policies selected"}), 400
        results = disable_policies(device, policies)
        return jsonify({"results": results})

    @app.post("/api/remediation/export/postman")
    def remediation_export_postman():
        gate = _license_required()
        if gate:
            return gate
        from app.services.remediation_service import export_postman
        from io import BytesIO
        device = app.config.get('remediation_device', {})
        payload = request.get_json(silent=True) or {}
        policies = payload.get("policies", [])
        json_str = export_postman(policies, device)
        buf = BytesIO(json_str.encode("utf-8"))
        today = _date.today().isoformat()
        return app.response_class(
            response=buf.read(),
            status=200,
            mimetype="application/json",
            headers={"Content-Disposition": f"attachment; filename=apo_remediation_{today}.postman_collection.json"},
        )

    @app.post("/api/remediation/export/csv")
    def remediation_export_csv():
        gate = _license_required()
        if gate:
            return gate
        from app.services.remediation_service import export_csv
        from io import BytesIO
        payload = request.get_json(silent=True) or {}
        csv_bytes = export_csv(
            payload.get("to_disable", []),
            payload.get("already_disabled", []),
        )
        today = _date.today().isoformat()
        return app.response_class(
            response=csv_bytes,
            status=200,
            mimetype="text/csv; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename=apo_remediation_{today}.csv"},
        )

    # ── License ────────────────────────────────────────────────────────────
    @app.get("/api/license/status")
    def license_status():
        info = get_license_info()
        if info:
            return jsonify({"licensed": True, "email": info.get("email"),
                            "issued": info.get("issued"),
                            # 기존 키에는 tier가 없다 → 단일 조직으로 간주(하위호환)
                            "tier": info.get("tier") or "single"})
        return jsonify({"licensed": False})

    # ── 화이트라벨 브랜딩 (MSP·컨설턴트 티어 전용) ─────────────────────────
    # 컨설턴트는 고객사에 제출하는 보고서에 자기 회사명을 실어야 한다.
    # 단일 조직 티어와의 실질 구분점이므로 서버측에서 티어를 강제한다.
    # frozen exe에서 DATA_DIR는 _MEIPASS(매 실행 새로 추출·종료 시 삭제) 하위라
    # 여기 쓰면 브랜딩이 재시작마다 사라진다(재비판 NEW-BUG). 라이선스 파일과
    # 같은 원칙으로 exe 인접 경로에 영속한다.
    if getattr(_sys, 'frozen', False):
        _BRANDING_FILE = Path(_sys.executable).parent / "branding.json"
    else:
        _BRANDING_FILE = DATA_DIR / "branding.json"

    def _load_branding() -> dict:
        try:
            data = json.loads(_BRANDING_FILE.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}

    def _branding_allowed() -> bool:
        info = get_license_info()
        return bool(info) and (info.get("tier") or "single") == "msp"

    @app.get("/api/report/branding")
    def report_branding_get():
        return jsonify({"allowed": _branding_allowed(),
                        "branding": _load_branding() if _branding_allowed() else {}})

    @app.post("/api/report/branding")
    def report_branding_set():
        if not _branding_allowed():
            return jsonify({"error": "Report branding requires an MSP / Consultant "
                            "license. Single-organization licenses export unbranded "
                            "reports."}), 403
        payload = request.get_json(silent=True) or {}
        # C0 제어문자는 openpyxl이 거부해 export 전체가 400으로 죽고(재비판
        # NEW-BUG), 개행·양방향 제어문자는 셀 표시를 왜곡한다. 전부 공백 치환.
        import re as _re
        _clean = lambda v: _re.sub(
            "[\\x00-\\x1f\\x7f\\u200e\\u200f\\u202a-\\u202e\\u2066-\\u2069]",
            " ", str(v))[:80].strip()
        branding = {
            "company":      _clean(payload.get("company", "")),
            "prepared_for": _clean(payload.get("prepared_for", "")),
        }
        _BRANDING_FILE.parent.mkdir(parents=True, exist_ok=True)
        _BRANDING_FILE.write_text(json.dumps(branding, ensure_ascii=False),
                                  encoding="utf-8")
        return jsonify({"ok": True, "branding": branding})

    @app.post("/api/license/activate")
    def license_activate():
        payload = request.get_json(silent=True) or {}
        key = str(payload.get("key", "")).strip()
        if not key:
            return jsonify({"error": "key is required"}), 400
        try:
            data = activate(key)
            return jsonify({"ok": True, "email": data.get("email"), "issued": data.get("issued")})
        except ValueError as e:
            return jsonify({"error": str(e)}), 400

    return app
