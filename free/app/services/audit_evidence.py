"""감사 증적(Review Evidence) 리포트 — B-06.

배경: PCI-DSS v4.0 Req 1.2.7은 조사한 표준 중 유일하게 숫자를 명시한
요구사항이다 — "Configurations of NSCs are reviewed at least once every
six months". 그런데 검토를 실제로 했다는 '증거'를 만드는 일은 수작업이다.
이 모듈은 config 스냅샷 시계열(수집은 Oxidized·FortiManager·수동 백업 등
무엇이든)을 받아, 검토 회의에 그대로 제출할 수 있는 증적 문서를 만든다.
수집은 남이 하고, 설명은 APO가 한다.

증적이 스스로 증명해야 하는 것:
  1. 무엇을 검토했나 — 스냅샷별 파일명 + SHA-256 (문서-파일 대응 성립)
  2. 그 사이 무엇이 바뀌었나 — 구간별 변경 요약 (정책 추가/삭제/수정, 객체)
  3. 누가 언제 승인했나 — 검토자 서명란 (빈 칸으로 제공)

시간 순서에 대한 정직한 한계: FortiGate config 파일에는 백업 시점이
기록되지 않는다. 스냅샷 순서는 호출자가 준 순서(오래된 것부터)를 그대로
믿으며, 문서에 그 사실을 명시한다.
"""
from __future__ import annotations

import hashlib
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .workbook_exporter import HEADER_FILL, HEADER_FONT, write_text_cell

# 인용은 구현이 있는 지금만 허용된다(backlog B-06 마케팅 규칙). 조항 원문을
# 그대로 싣는 이유: 검토 회의에서 "왜 6개월이냐"는 질문에 문서가 답해야 한다.
PCI_DSS_CITATION = (
    "PCI-DSS v4.0 Requirement 1.2.7: \"Configurations of NSCs [network "
    "security controls] are reviewed at least once every six months to "
    "confirm they are relevant and effective.\""
)

ORDER_NOTE = (
    "Snapshot order: FortiGate configuration files do not record when the "
    "backup was taken, so snapshots are listed in the order they were "
    "provided (oldest first, as stated by the operator). Each snapshot's "
    "SHA-256 lets you tie this document back to the exact files reviewed."
)


def snapshot_meta(filename: str, raw: bytes, parsed: dict) -> dict:
    """스냅샷 식별 정보. SHA-256이 문서-파일 대응의 근거다."""
    return {
        "filename": filename,
        "sha256": hashlib.sha256(raw).hexdigest(),
        "hostname": (parsed.get("meta") or {}).get("hostname", ""),
        "config_version": (parsed.get("meta") or {}).get("config_version", ""),
        "policy_count": len(parsed.get("firewall_policy") or []),
    }


def interval_summary(diff: dict) -> dict:
    """_compute_config_diff 결과에서 증적에 필요한 요약만 추린다."""
    s = diff.get("summary") or {}
    changed_ids = [str(c.get("policy_id")) for c in (diff.get("changed_policies") or [])]
    added_ids = [str(p.get("policy_id")) for p in (diff.get("added_policies") or [])]
    removed_ids = [str(p.get("policy_id")) for p in (diff.get("removed_policies") or [])]
    return {
        "added_policies": s.get("added_policies", 0),
        "removed_policies": s.get("removed_policies", 0),
        "changed_policies": s.get("changed_policies", 0),
        "added_objects": s.get("added_objects", 0),
        "removed_objects": s.get("removed_objects", 0),
        "other_changes": s.get("other_changes", 0),
        "added_ids": added_ids[:50],
        "removed_ids": removed_ids[:50],
        "changed_ids": changed_ids[:50],
    }


def build_evidence_workbook(timeline: dict, report_meta: dict | None = None) -> bytes:
    """timeline: {"snapshots": [snapshot_meta...], "intervals": [{from,to,summary}]}"""
    wb = Workbook()
    wb.remove(wb.active)
    rm = report_meta or {}

    # ── 1. Review Evidence (표지 + 서명란) ─────────────────────────────────
    ws = wb.create_sheet("Review Evidence")
    ws.sheet_view.showGridLines = False
    r = 1
    ws.cell(row=r, column=1, value="Firewall Ruleset Review — Evidence Record").font = \
        Font(bold=True, size=14)
    r += 1
    branding = rm.get("branding") or {}
    if branding.get("company"):
        line = f"Prepared by {branding['company']}"
        if branding.get("prepared_for"):
            line += f" for {branding['prepared_for']}"
        write_text_cell(ws, r, 1, line).font = Font(bold=True, size=11)
        r += 1
    r += 1

    snaps = timeline.get("snapshots") or []
    hosts = sorted({s.get("hostname") for s in snaps if s.get("hostname")})
    kv = [
        ("Generated", rm.get("generated_at", "")),
        ("Tool version", rm.get("apo_version", "")),
        ("Device(s)", ", ".join(hosts)),
        ("Snapshots reviewed", str(len(snaps))),
        ("Review basis", PCI_DSS_CITATION),
    ]
    for k, v in kv:
        write_text_cell(ws, r, 1, k).font = Font(bold=True)
        c = write_text_cell(ws, r, 2, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1

    write_text_cell(ws, r, 1, ORDER_NOTE).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 2

    # 서명란 — 검토는 사람이 하고, 문서는 그 사실을 담는다
    ws.cell(row=r, column=1, value="Review sign-off").font = Font(bold=True, size=12)
    r += 1
    for col, h in enumerate(("Reviewer (name / role)", "Review date",
                             "Decision (approved / changes required)",
                             "Notes"), 1):
        cell = ws.cell(row=r, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    r += 1
    for _ in range(3):        # 검토자 복수 서명 여지
        for col in range(1, 5):
            ws.cell(row=r, column=col, value="")
        r += 1
    for col, w in (("A", 30), ("B", 90), ("C", 34), ("D", 30)):
        ws.column_dimensions[col].width = w

    # ── 2. Snapshots (식별 정보) ───────────────────────────────────────────
    ws2 = wb.create_sheet("Snapshots")
    headers = ["#", "Filename", "SHA-256", "Hostname", "Config version",
               "Policy count"]
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for i, s in enumerate(snaps, 1):
        ws2.cell(row=i + 1, column=1, value=i)
        write_text_cell(ws2, i + 1, 2, str(s.get("filename", "")))
        write_text_cell(ws2, i + 1, 3, str(s.get("sha256", "")))
        write_text_cell(ws2, i + 1, 4, str(s.get("hostname", "")))
        write_text_cell(ws2, i + 1, 5, str(s.get("config_version", "")))
        ws2.cell(row=i + 1, column=6, value=s.get("policy_count", 0))
    for col, w in (("A", 5), ("B", 40), ("C", 66), ("D", 20), ("E", 40), ("F", 12)):
        ws2.column_dimensions[col].width = w

    # ── 3. Changes (구간별 변경) ───────────────────────────────────────────
    ws3 = wb.create_sheet("Changes")
    headers = ["Interval", "From", "To", "Policies added", "Policies removed",
               "Policies changed", "Objects ±", "Other sections",
               "Added IDs", "Removed IDs", "Changed IDs"]
    for col, h in enumerate(headers, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    r = 2
    for i, iv in enumerate(timeline.get("intervals") or [], 1):
        sm = iv.get("summary") or {}
        ws3.cell(row=r, column=1, value=i)
        write_text_cell(ws3, r, 2, str(iv.get("from", "")))
        write_text_cell(ws3, r, 3, str(iv.get("to", "")))
        ws3.cell(row=r, column=4, value=sm.get("added_policies", 0))
        ws3.cell(row=r, column=5, value=sm.get("removed_policies", 0))
        ws3.cell(row=r, column=6, value=sm.get("changed_policies", 0))
        write_text_cell(ws3, r, 7,
                        f"+{sm.get('added_objects', 0)} / -{sm.get('removed_objects', 0)}")
        ws3.cell(row=r, column=8, value=sm.get("other_changes", 0))
        write_text_cell(ws3, r, 9, ", ".join(sm.get("added_ids") or []))
        write_text_cell(ws3, r, 10, ", ".join(sm.get("removed_ids") or []))
        write_text_cell(ws3, r, 11, ", ".join(sm.get("changed_ids") or []))
        for cell in ws3[r]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1
    for col, w in (("A", 8), ("B", 34), ("C", 34), ("D", 9), ("E", 9), ("F", 9),
                   ("G", 12), ("H", 9), ("I", 30), ("J", 30), ("K", 30)):
        ws3.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
